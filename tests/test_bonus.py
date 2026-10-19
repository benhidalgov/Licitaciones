import io
import os
import unittest
from datetime import datetime

import openpyxl

import app as flask_app
import database


def _lic(**extra):
    """Licitacion minima valida para create_licitacion."""
    base = {
        "id_licitacion": "100-1-LP26",
        "titulo": "Servicio de hospedaje institucional",
        "organismo": "Ministerio de Educacion",
        "categoria": "Alojamiento",
        "modalidad": "Licitacion Publica",
        "region": "Region Metropolitana",
        "estado_embudo": "Identificada",
        "validacion_adjuntos": "Valida",
        "tiene_anexo4": 1, "tiene_escrituras": 1,
        "tiene_poderes": 1, "tiene_vigencias": 1,
        "presupuesto_mandante": 10000000,
        "costo_base_hotel": 7000000,
        "fecha_publicacion": "2026-10-01",
        "fecha_cierre": "2026-11-30",
        "descripcion_tdr": "Alojamiento para funcionarios",
    }
    base.update(extra)
    return base


class TestBonus(unittest.TestCase):
    def setUp(self):
        self.db_name = "test_bonus_licitaciones.db"
        database.DB_NAME = self.db_name
        if os.path.exists(self.db_name):
            try:
                os.remove(self.db_name)
            except PermissionError:
                pass
        database.init_db()
        self._vaciar_licitaciones()

        flask_app.app.config["TESTING"] = True
        flask_app.app.config["CSRF_ENABLED"] = False
        self.client = flask_app.app.test_client()
        with self.client.session_transaction() as sess:
            sess["user"] = "daniela"
            sess["nombre"] = "Daniela"

    @staticmethod
    def _vaciar_licitaciones():
        """init_db siembra licitaciones de ejemplo; estos tests necesitan una base limpia."""
        conn = database.get_db_connection()
        try:
            conn.execute("DELETE FROM licitaciones")
            conn.execute("DELETE FROM historial_oc")
            conn.commit()
        finally:
            conn.close()

    def tearDown(self):
        if os.path.exists(self.db_name):
            try:
                os.remove(self.db_name)
            except PermissionError:
                pass

    # -- Bonus 1: monitoreo de clientes clave --------------------------------

    def test_detecta_clientes_clave_con_y_sin_acentos(self):
        casos = [
            ("Ministerio de Educación", "Mineduc"),
            ("Subsecretaría de Educación", "Mineduc"),
            ("Ministerio de Salud", "Minsal"),
            ("Universidad de Chile", "U. de Chile"),
            ("Universidad de Santiago de Chile", "Usach"),
            ("Productoras Audiovisuales SpA", "Productoras"),
            ("SEREMI de Turismo", ""),
        ]
        for texto, esperado in casos:
            self.assertEqual(
                database.detectar_cliente_clave(texto), esperado, msg=texto
            )

    def test_filtro_por_cliente_clave_acota_el_listado(self):
        database.create_licitacion(_lic(id_licitacion="1", organismo="Ministerio de Salud"))
        database.create_licitacion(_lic(id_licitacion="2", organismo="Dirección de Compras"))
        database.create_licitacion(_lic(id_licitacion="3", organismo="Universidad de Chile"))

        self.assertEqual(database.count_licitaciones(cliente_clave="Minsal"), 1)
        self.assertEqual(database.count_licitaciones(cliente_clave="U. de Chile"), 1)
        self.assertEqual(len(database.get_all_licitaciones(cliente_clave="Minsal")), 1)
        # Cliente fuera de la lista configurable: no devuelve nada en vez de todo.
        self.assertEqual(database.count_licitaciones(cliente_clave="Nadie"), 0)
        # Sin filtro sigue trayendo todo.
        self.assertEqual(database.count_licitaciones(), 3)

    def test_tablero_filtra_por_cliente_clave(self):
        database.create_licitacion(_lic(id_licitacion="1", organismo="Ministerio de Salud",
                                        titulo="Hospedaje Minsal"))
        database.create_licitacion(_lic(id_licitacion="2", organismo="Dirección de Compras",
                                        titulo="Hospedaje Generico"))

        resp = self.client.get("/licitaciones?cliente_clave=Minsal")
        html = resp.get_data(as_text=True)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Hospedaje Minsal", html)
        self.assertNotIn("Hospedaje Generico", html)
        self.assertIn("Cliente Clave", html)

    def test_tablero_muestra_etiqueta_de_cliente_clave(self):
        database.create_licitacion(_lic(id_licitacion="1", organismo="Ministerio de Salud"))
        html = self.client.get("/licitaciones").get_data(as_text=True)
        self.assertIn("Cliente: Minsal", html)

    # -- Bonus 2: benchmarking de adjudicaciones -----------------------------

    def test_guardar_benchmarking_y_calcular_desviacion(self):
        database.create_licitacion(_lic(id_licitacion="1"))
        resp = self.client.post(
            "/licitaciones/1/benchmarking",
            data={
                "adjudicatario": "Hotel Competidor SpA",
                "puntaje_adjudicacion": "78,5",
                "precio_final": "$12.345.678",
            },
        )
        self.assertEqual(resp.status_code, 302)

        lic = database.get_licitacion_by_id("1")
        self.assertEqual(lic["adjudicatario"], "Hotel Competidor SpA")
        self.assertEqual(lic["puntaje_adjudicacion"], 78.5)
        self.assertEqual(lic["precio_final"], 12345678)

        benchmark = database.get_benchmarking()
        self.assertEqual(benchmark["adjudicaciones"], 1)
        fila = benchmark["filas"][0]
        # 12.345.678 vs presupuesto 10.000.000 -> +23.5% observado
        self.assertAlmostEqual(fila["desviacion_pct"], 23.5, places=1)
        # margen sobre nuestro costo de 7.000.000
        self.assertEqual(fila["margen_vs_costo"], 12345678 - 7000000)
        self.assertEqual(benchmark["precio_promedio"], 12345678)

    def test_benchmarking_rechaza_puntaje_fuera_de_rango(self):
        database.create_licitacion(_lic(id_licitacion="1"))
        self.client.post(
            "/licitaciones/1/benchmarking",
            data={"adjudicatario": "X", "puntaje_adjudicacion": "150", "precio_final": "1000"},
        )
        lic = database.get_licitacion_by_id("1")
        self.assertEqual(lic["adjudicatario"], "")
        self.assertEqual(lic["precio_final"], 0)

    def test_analitica_renderiza_vacia(self):
        resp = self.client.get("/analitica")
        html = resp.get_data(as_text=True)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Benchmarking de Adjudicaciones", html)
        self.assertIn("Historial de Ordenes de Compra", html)
        self.assertIn("Todavia no hay adjudicaciones registradas", html)

    # -- Bonus 3: historial de OC --------------------------------------------

    def _hoja_excel(self, encabezados, filas):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(encabezados)
        for fila in filas:
            hoja.append(fila)
        buffer = io.BytesIO()
        libro.save(buffer)
        buffer.seek(0)
        return buffer

    def test_importa_excel_de_oc_con_formato_chileno(self):
        datos = self._hoja_excel(
            ["Fecha O.C.", "Cliente", "Descripción", "Monto Total", "Modalidad"],
            [
                ["15/03/2026", "Ministerio de Educación", "Alojamiento delegados", "$1.234.567", "Licitación Privada"],
                [datetime(2026, 4, 2), "Minsal", "Coffee break jornada", "2500000", "Compra Ágil"],
            ],
        )
        resp = self.client.post(
            "/analitica/historial",
            data={"archivo": (datos, "historial.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 302)

        resumen = database.get_resumen_historial_oc()
        self.assertEqual(resumen["registros"], 2)
        self.assertEqual(resumen["monto_total"], 1234567 + 2500000)

        filas = database.get_historial_oc()
        por_organismo = {f["organismo"]: f for f in filas}
        self.assertEqual(por_organismo["Ministerio de Educación"]["fecha"], "2026-03-15")
        self.assertEqual(por_organismo["Ministerio de Educación"]["monto"], 1234567)
        self.assertEqual(por_organismo["Minsal"]["fecha"], "2026-04-02")
        self.assertEqual(por_organismo["Minsal"]["monto"], 2500000)

    def test_importa_csv_cuando_no_es_excel(self):
        csv_texto = (
            "organismo,objeto,monto\n"
            "Universidad de Chile,Arriendo de salon,900000\n"
        )
        resp = self.client.post(
            "/analitica/historial",
            data={"archivo": (io.BytesIO(csv_texto.encode("utf-8-sig")), "oc.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(database.get_resumen_historial_oc()["registros"], 1)

    def test_rechaza_archivo_sin_columnas_reconocibles(self):
        datos = self._hoja_excel(["Serie", "Correlativo"], [["A", 500]])
        filas = [
            list(r) for r in openpyxl.load_workbook(
                io.BytesIO(datos.getvalue()), read_only=True
            ).active.iter_rows(values_only=True)
        ]
        self.assertEqual(database.importar_historial_oc(filas), 0)

    def test_rechaza_formato_no_soportado(self):
        resp = self.client.post(
            "/analitica/historial",
            data={"archivo": (io.BytesIO(b"datos"), "archivo.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(database.get_resumen_historial_oc()["registros"], 0)

    def test_cruce_de_experiencia_por_organismo(self):
        database.importar_historial_oc(
            [
                ["fecha", "organismo", "objeto", "monto"],
                ["2025-05-01", "Ministerio de Educación", "Hospedaje", 800000],
                ["2025-06-01", "Ministerio de Educación", "Salón", 500000],
                ["2025-07-01", "Minsal", "Catering", 300000],
            ],
            archivo="prueba.csv",
        )
        experiencia = database.get_experiencia_organismo("Ministerio de Educación")
        self.assertEqual(experiencia["registros"], 2)
        self.assertEqual(experiencia["monto_total"], 1300000)

        ninguna = database.get_experiencia_organismo("")
        self.assertEqual(ninguna["registros"], 0)

    def test_detalle_muestra_experiencia_y_cliente_clave(self):
        database.create_licitacion(_lic(id_licitacion="1", organismo="Ministerio de Educación"))
        database.importar_historial_oc(
            [
                ["fecha", "organismo", "objeto", "monto"],
                ["2025-05-01", "Ministerio de Educación", "Hospedaje", 800000],
            ],
            archivo="prueba.csv",
        )
        html = self.client.get("/licitaciones/1").get_data(as_text=True)
        self.assertEqual(html.count("1 OC"), 1)
        self.assertIn("Cliente clave: Mineduc", html)
        self.assertIn("Benchmarking de Adjudicacion", html)


if __name__ == "__main__":
    unittest.main()
